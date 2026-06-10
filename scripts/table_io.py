"""Unified tabular reader for csv / tsv / xlsx / numbers.

read_table(path) -> (headers, rows), where rows are dicts keyed by header.
All values are stripped strings ('' for empty cells), so downstream callers
see one consistent shape regardless of source format.
"""
from __future__ import annotations

import csv
from pathlib import Path


def read_table(path: str | Path) -> tuple[list[str], list[dict]]:
    p = str(path).lower()
    if p.endswith(('.xlsx', '.xls', '.xlsm')):
        return _read_xlsx(path)
    if p.endswith('.numbers'):
        return _read_numbers(path)
    delim = '\t' if p.endswith('.tsv') else ','
    return _read_csv(path, delim)


def _norm(v) -> str:
    return str(v).strip() if v is not None else ''


def _rows_from_grid(headers: list[str], data_rows) -> list[dict]:
    out = []
    for raw in data_rows:
        if all(v is None for v in raw):
            continue
        out.append({h: _norm(v) for h, v in zip(headers, raw)})
    return out


def _read_csv(path, delim: str) -> tuple[list[str], list[dict]]:
    with open(path, newline='', encoding='utf-8-sig') as fh:
        reader = csv.DictReader(fh, delimiter=delim)
        headers = list(reader.fieldnames or [])
        rows = [{h: _norm(row.get(h)) for h in headers} for row in reader]
    return headers, rows


def _read_xlsx(path) -> tuple[list[str], list[dict]]:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    grid = list(ws.iter_rows(values_only=True))
    if not grid:
        return [], []
    headers = [_norm(h) for h in grid[0]]
    return headers, _rows_from_grid(headers, grid[1:])


def _read_numbers(path) -> tuple[list[str], list[dict]]:
    from numbers_parser import Document
    doc = Document(path)
    sheets = {s.name: s for s in doc.sheets}
    sheet = sheets.get('HELM_Builder') or doc.sheets[0]
    grid = [[c.value for c in row] for row in sheet.tables[0].rows()]
    if not grid:
        return [], []
    headers = [_norm(h) for h in grid[0]]
    return headers, _rows_from_grid(headers, grid[1:])
