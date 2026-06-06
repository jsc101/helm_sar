#!/usr/bin/env python3
"""
run_sar.py — One-command SAR report for a HELM peptide library.

Reads a CSV, Excel, or Apple Numbers file containing compound names, HELM strings,
and optional activity data. Generates an interactive HTML report with:
  - NW sequence alignment coloured by chemistry (Zappo) and charge/LogP
  - Sortable columns: alignment score, % identity, MW, activity, changes vs ref
  - Two groups: scaffold analogues (Group A) and extended library (Group B)

Usage
-----
    python run_sar.py --input library.csv
    python run_sar.py --input library.xlsx --activity IC50_nM --ref "Parent"
    python run_sar.py --input library.numbers --out report.html --no-open

Input format (CSV / Excel)
--------------------------
Required columns (case-insensitive match):
    Name   — compound identifier
    HELM   — HELM V2 string

Optional columns (any name, specify with --activity):
    IC50_nM, pIC50, %inhibition, Ki_nM, ...

The reference compound is auto-detected as the row whose Name contains
"parent", "ref", "wt", "wildtype", or "reference" (case-insensitive).
If none matches, the first row is used.

Example CSV
-----------
    Name,HELM,IC50_nM
    Parent,"PEPTIDE1{A.K.G.F}$PEPTIDE1,PEPTIDE1,1:R1-4:R2$$$V2.0",45.0
    Analog_1,"PEPTIDE1{A.R.G.F}$PEPTIDE1,PEPTIDE1,1:R1-4:R2$$$V2.0",12.0
"""

from __future__ import annotations

import sys
import os
import argparse
import logging
import subprocess
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
logging.disable(logging.WARNING)

from monomer_db.monomer_db import MonomerDB
from scripts.report import build_data, build_html

_CYCPEPT_JSON = Path(__file__).parent / 'monomer_db' / 'cycpeptmpdb_monomers.json'


# ── Input readers ────────────────────────────────────────────────────────────────

def _ref_keywords():
    return ('parent', 'ref', 'wt', 'wildtype', 'reference')


def _is_ref_name(name: str) -> bool:
    n = name.lower()
    return any(k in n for k in _ref_keywords())


def read_csv(path: str, activity_col: str | None) -> tuple[list[tuple], dict]:
    import csv
    pairs, activity = [], {}
    with open(path, newline='', encoding='utf-8-sig') as fh:
        reader = csv.DictReader(fh)
        headers = reader.fieldnames or []
        name_col = _find_col(headers, ['name', 'id', 'compound', 'cmpd'])
        helm_col = _find_col(headers, ['helm', 'helm_string', 'sequence'])
        act_col  = _find_col(headers, [activity_col] if activity_col else []) if activity_col else None
        for row in reader:
            name = row.get(name_col, '').strip()
            helm = row.get(helm_col, '').strip()
            if helm and 'PEPTIDE' in helm:
                pairs.append((name, helm))
                if act_col and act_col in row:
                    v = row[act_col].strip()
                    try:
                        activity[name] = float(v)
                    except ValueError:
                        activity[name] = v or None
    return pairs, activity


def read_excel(path: str, activity_col: str | None) -> tuple[list[tuple], dict]:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], {}
    headers = [str(h).strip() if h else '' for h in rows[0]]
    name_col = _find_col_idx(headers, ['name', 'id', 'compound', 'cmpd'])
    helm_col = _find_col_idx(headers, ['helm', 'helm_string', 'sequence'])
    act_col  = _find_col_idx(headers, [activity_col] if activity_col else []) if activity_col else None
    pairs, activity = [], {}
    for row in rows[1:]:
        name = str(row[name_col]).strip() if name_col is not None and row[name_col] else ''
        helm = str(row[helm_col]).strip() if helm_col is not None and row[helm_col] else ''
        if helm and 'PEPTIDE' in helm:
            pairs.append((name, helm))
            if act_col is not None and act_col < len(row) and row[act_col] is not None:
                v = row[act_col]
                try:
                    activity[name] = float(v)
                except (TypeError, ValueError):
                    activity[name] = str(v) if v else None
    return pairs, activity


def read_numbers(path: str, activity_col: str | None) -> tuple[list[tuple], dict]:
    from numbers_parser import Document
    doc = Document(path)
    sheets = {s.name: s for s in doc.sheets}
    # prefer HELM_Builder sheet; fall back to first sheet
    sheet = sheets.get('HELM_Builder') or doc.sheets[0]
    table = sheet.tables[0]
    all_rows = list(table.rows())
    if not all_rows:
        return [], {}
    headers = [str(c.value).strip() if c.value else '' for c in all_rows[0]]
    name_col = _find_col_idx(headers, ['name', 'id', 'compound', 'cmpd'])
    helm_col = _find_col_idx(headers, ['helm', 'helm_string', 'sequence'])
    act_col  = _find_col_idx(headers, [activity_col] if activity_col else []) if activity_col else None
    pairs, activity = [], {}
    for row in all_rows[1:]:
        name = str(row[name_col].value).strip() if name_col is not None and row[name_col].value else ''
        helm = str(row[helm_col].value).strip() if helm_col is not None and len(row) > helm_col and row[helm_col].value else ''
        if helm and 'PEPTIDE' in helm:
            pairs.append((name, helm))
            if act_col is not None and act_col < len(row) and row[act_col].value is not None:
                v = row[act_col].value
                try:
                    activity[name] = float(v)
                except (TypeError, ValueError):
                    activity[name] = str(v) if v else None
    return pairs, activity


def _find_col(headers: list[str], candidates: list[str]) -> str | None:
    hl = [h.lower() for h in headers]
    for c in candidates:
        for i, h in enumerate(hl):
            if c.lower() in h:
                return headers[i]
    return None


def _find_col_idx(headers: list[str], candidates: list[str]) -> int | None:
    hl = [h.lower() for h in headers]
    for c in candidates:
        for i, h in enumerate(hl):
            if c.lower() in h:
                return i
    return None


def read_input(path: str, activity_col: str | None) -> tuple[list[tuple], dict]:
    p = path.lower()
    if p.endswith('.csv') or p.endswith('.tsv'):
        return read_csv(path, activity_col)
    if p.endswith('.xlsx') or p.endswith('.xls'):
        return read_excel(path, activity_col)
    if p.endswith('.numbers'):
        return read_numbers(path, activity_col)
    # try CSV as default
    return read_csv(path, activity_col)


# ── Main ─────────────────────────────────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser(
        description='Generate a HELM peptide SAR report (HTML).',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    ap.add_argument('--input', '-i', required=True,
                    help='Input file: CSV, Excel (.xlsx), or Apple Numbers (.numbers)')
    ap.add_argument('--out', '-o', default='sar_report.html',
                    help='Output HTML path (default: sar_report.html)')
    ap.add_argument('--activity', '-a', default=None,
                    help='Column name for activity data (e.g. IC50_nM, pIC50)')
    ap.add_argument('--ref', '-r', default=None,
                    help='Exact name of reference compound (default: auto-detect)')
    ap.add_argument('--max-rows', type=int, default=50,
                    help='Max compounds shown in table (default: 50)')
    ap.add_argument('--min-identity', type=float, default=0.3,
                    help='Exclude compounds with identity below this fraction (default: 0.3)')
    ap.add_argument('--no-open', action='store_true',
                    help='Do not open the HTML in the browser after generation')
    args = ap.parse_args()

    print('Loading MonomerDB…')
    extra = [str(_CYCPEPT_JSON)] if _CYCPEPT_JSON.exists() else []
    db = MonomerDB(extra_sources=extra)

    print(f'Reading {args.input}…')
    pairs, activity_map = read_input(args.input, args.activity)
    if not pairs:
        print('ERROR: no HELM entries found. Check column names (Name, HELM).', file=sys.stderr)
        sys.exit(1)
    print(f'  {len(pairs)} HELM entries')
    if activity_map:
        print(f'  {len(activity_map)} activity values ({args.activity})')

    # Override reference if explicitly named
    if args.ref:
        pairs = sorted(pairs, key=lambda p: 0 if p[0] == args.ref else 1)

    print('Parsing and aligning…')
    ref_obj, rows = build_data(pairs, db)

    for r in rows:
        status = 'REF' if r['is_ref'] else (f"score={r['score']:.3f}" if r['score'] is not None else 'diff-len')
        mw_str = f"MW={r['mw']:.1f}" if r['mw'] is not None else 'MW=?'
        act_str = f"  act={activity_map[r['id']]}" if r['id'] in activity_map else ''
        print(f"  {r['id']:20s}  {mw_str:14s}  {status}{act_str}")

    print('Building HTML…')
    act_label = args.activity or 'Activity'
    out_html = build_html(ref_obj, rows, db=db,
                          activity_map=activity_map or None,
                          act_label=act_label,
                          max_rows=args.max_rows,
                          min_identity=args.min_identity)

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(out_html, encoding='utf-8')
    print(f'Written → {out_path}  ({out_path.stat().st_size // 1024} KB)')

    if not args.no_open:
        subprocess.Popen(['open', str(out_path)])


if __name__ == '__main__':
    main()
